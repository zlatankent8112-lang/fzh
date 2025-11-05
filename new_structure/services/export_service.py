"""
Export Service for Fee Management System
Handles Excel and PDF export functionality
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from io import BytesIO
from datetime import datetime


class ExportService:
    """Service for exporting data to Excel and PDF"""
    
    @staticmethod
    def export_analytics_to_excel(data):
        """Export analytics data to Excel"""
        wb = Workbook()
        
        # Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Summary"
        
        # Header
        ws_summary['A1'] = "Fee Collection Analytics Report"
        ws_summary['A1'].font = Font(size=16, bold=True)
        ws_summary['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        
        # Summary Stats
        ws_summary['A4'] = "Expected Revenue"
        ws_summary['B4'] = float(data['total_expected'])
        ws_summary['A5'] = "Collected"
        ws_summary['B5'] = float(data['total_collected'])
        ws_summary['A6'] = "Outstanding"
        ws_summary['B6'] = float(data['total_outstanding'])
        ws_summary['A7'] = "Collection Rate"
        ws_summary['B7'] = f"{data['collection_rate']:.1f}%"
        
        # Style summary
        for row in range(4, 8):
            ws_summary[f'A{row}'].font = Font(bold=True)
            ws_summary[f'B{row}'].number_format = '#,##0.00'
        
        # Grade Stats Sheet
        ws_grades = wb.create_sheet("By Grade")
        ws_grades['A1'] = "Collection by Grade"
        ws_grades['A1'].font = Font(size=14, bold=True)
        
        headers = ['Grade', 'Students', 'Expected (KES)', 'Collected (KES)', 'Outstanding (KES)', 'Rate (%)']
        for col, header in enumerate(headers, 1):
            cell = ws_grades.cell(3, col, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        row = 4
        for grade_data in data['grade_stats']:
            grade_name, student_count, expected, collected, outstanding = grade_data
            rate = (collected / expected * 100) if expected > 0 else 0
            ws_grades.cell(row, 1, grade_name)
            ws_grades.cell(row, 2, student_count)
            ws_grades.cell(row, 3, float(expected))
            ws_grades.cell(row, 4, float(collected))
            ws_grades.cell(row, 5, float(outstanding))
            ws_grades.cell(row, 6, f"{rate:.1f}%")
            row += 1
        
        # Format numbers
        for row in range(4, row):
            for col in range(3, 6):
                ws_grades.cell(row, col).number_format = '#,##0.00'
        
        # Adjust column widths
        for ws in [ws_summary, ws_grades]:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[column_letter].width = max_length + 2
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_defaulters_to_pdf(defaulters):
        """Export defaulter list to PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#333333'),
            spaceAfter=30,
        )
        elements.append(Paragraph("Fee Defaulters Report", title_style))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Table data
        data = [['Student', 'Admission No', 'Grade', 'Total Fees', 'Paid', 'Outstanding']]
        
        for student, total_fees, total_paid, balance in defaulters:
            data.append([
                student.name[:25],
                student.admission_number or 'N/A',
                student.grade.name if student.grade else 'N/A',
                f"KES {float(total_fees):,.0f}",
                f"KES {float(total_paid):,.0f}",
                f"KES {float(balance):,.0f}"
            ])
        
        # Create table
        table = Table(data, colWidths=[2*inch, 1.2*inch, 1*inch, 1.2*inch, 1.2*inch, 1.2*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f0f0')]),
        ]))
        
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def export_balance_report_to_excel(students_with_balances):
        """Export balance report to Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Balance Report"
        
        # Header
        ws['A1'] = "Student Fee Balance Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        
        # Column headers
        headers = ['Student Name', 'Admission No', 'Grade', 'Stream', 'Total Fees', 'Amount Paid', 'Balance', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(4, col, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        # Data rows
        row = 5
        for student_data in students_with_balances:
            ws.cell(row, 1, student_data['name'])
            ws.cell(row, 2, student_data['admission_number'])
            ws.cell(row, 3, student_data['grade'])
            ws.cell(row, 4, student_data['stream'])
            ws.cell(row, 5, float(student_data['total_fees']))
            ws.cell(row, 6, float(student_data['amount_paid']))
            ws.cell(row, 7, float(student_data['balance']))
            ws.cell(row, 8, student_data['status'])
            row += 1
        
        # Format currency columns
        for r in range(5, row):
            for col in range(5, 8):
                ws.cell(r, col).number_format = '#,##0.00'
        
        # Adjust widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = max_length + 2
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
