"""
Fee Export Service - Excel and PDF export functionality
Handles exporting fee data, analytics, and reports to Excel and PDF formats
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfgen import canvas
from datetime import datetime
from decimal import Decimal
from io import BytesIO


class FeeExportService:
    """Service for exporting fee management data"""
    
    @staticmethod
    def export_analytics_to_excel(analytics_data):
        """
        Export analytics dashboard data to Excel
        
        Args:
            analytics_data: Dictionary containing analytics data
            
        Returns:
            BytesIO: Excel file as bytes
        """
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # 1. Summary Sheet
        ws_summary = wb.create_sheet("Summary", 0)
        FeeExportService._create_summary_sheet(ws_summary, analytics_data)
        
        # 2. Grade Breakdown Sheet
        ws_grades = wb.create_sheet("Grade Breakdown", 1)
        FeeExportService._create_grade_sheet(ws_grades, analytics_data.get('grade_stats', []))
        
        # 3. Stream Breakdown Sheet
        ws_streams = wb.create_sheet("Stream Breakdown", 2)
        FeeExportService._create_stream_sheet(ws_streams, analytics_data.get('stream_stats', []))
        
        # 4. Payment Methods Sheet
        ws_methods = wb.create_sheet("Payment Methods", 3)
        FeeExportService._create_payment_methods_sheet(ws_methods, analytics_data.get('payment_method_stats', []))
        
        # 5. Defaulters Sheet
        ws_defaulters = wb.create_sheet("Defaulters", 4)
        FeeExportService._create_defaulters_sheet(ws_defaulters, analytics_data.get('defaulters', []))
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    @staticmethod
    def _create_summary_sheet(ws, data):
        """Create summary statistics sheet"""
        # Title
        ws['A1'] = 'Fee Collection Analytics Summary'
        ws['A1'].font = Font(size=16, bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.merge_cells('A1:D1')
        
        # Generation date
        ws['A2'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}'
        ws['A2'].alignment = Alignment(horizontal='center')
        ws.merge_cells('A2:D2')
        
        # Headers
        headers = ['Metric', 'Value', 'Currency', 'Notes']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Data rows
        metrics = [
            ('Expected Revenue', data.get('total_expected', 0), 'KES', 'Total fees charged'),
            ('Collected', data.get('total_collected', 0), 'KES', 'Total payments received'),
            ('Outstanding', data.get('total_outstanding', 0), 'KES', 'Pending payments'),
            ('Collection Rate', data.get('collection_rate', 0), '%', 'Collection efficiency'),
            ('', '', '', ''),
            ('Fully Paid Students', data.get('fully_paid_count', 0), 'count', 'Students with zero balance'),
            ('Partially Paid Students', data.get('partially_paid_count', 0), 'count', 'Students with some payment'),
            ('Not Paid Students', data.get('not_paid_count', 0), 'count', 'Students with no payment'),
        ]
        
        row = 5
        for metric, value, currency, notes in metrics:
            ws.cell(row=row, column=1, value=metric).font = Font(bold=True if metric else False)
            if currency == 'KES':
                ws.cell(row=row, column=2, value=float(value)).number_format = '#,##0.00'
            elif currency == '%':
                ws.cell(row=row, column=2, value=float(value)).number_format = '0.00'
            else:
                ws.cell(row=row, column=2, value=value)
            ws.cell(row=row, column=3, value=currency)
            ws.cell(row=row, column=4, value=notes)
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 40
    
    @staticmethod
    def _create_grade_sheet(ws, grade_stats):
        """Create grade breakdown sheet"""
        # Title
        ws['A1'] = 'Collection by Grade'
        ws['A1'].font = Font(size=14, bold=True)
        ws.merge_cells('A1:G1')
        
        # Headers
        headers = ['Grade', 'Students', 'Expected (KES)', 'Collected (KES)', 'Outstanding (KES)', 'Rate (%)', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        
        # Data
        row = 4
        for grade_name, student_count, expected, collected, outstanding in grade_stats:
            rate = (float(collected) / float(expected) * 100) if expected > 0 else 0
            status = 'Excellent' if rate >= 80 else 'Good' if rate >= 50 else 'Poor'
            
            ws.cell(row=row, column=1, value=grade_name)
            ws.cell(row=row, column=2, value=student_count)
            ws.cell(row=row, column=3, value=float(expected)).number_format = '#,##0.00'
            ws.cell(row=row, column=4, value=float(collected)).number_format = '#,##0.00'
            ws.cell(row=row, column=5, value=float(outstanding)).number_format = '#,##0.00'
            ws.cell(row=row, column=6, value=rate).number_format = '0.00'
            
            status_cell = ws.cell(row=row, column=7, value=status)
            if rate >= 80:
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif rate >= 50:
                status_cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            else:
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            
            row += 1
        
        # Adjust widths
        for col in range(1, 8):
            ws.column_dimensions[get_column_letter(col)].width = 18
    
    @staticmethod
    def _create_stream_sheet(ws, stream_stats):
        """Create stream breakdown sheet"""
        ws['A1'] = 'Collection by Stream'
        ws['A1'].font = Font(size=14, bold=True)
        ws.merge_cells('A1:G1')
        
        headers = ['Stream', 'Grade', 'Students', 'Expected (KES)', 'Collected (KES)', 'Outstanding (KES)', 'Rate (%)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        
        row = 4
        for stream_name, grade_name, student_count, expected, collected, outstanding in stream_stats:
            rate = (float(collected) / float(expected) * 100) if expected > 0 else 0
            
            ws.cell(row=row, column=1, value=stream_name)
            ws.cell(row=row, column=2, value=grade_name)
            ws.cell(row=row, column=3, value=student_count)
            ws.cell(row=row, column=4, value=float(expected)).number_format = '#,##0.00'
            ws.cell(row=row, column=5, value=float(collected)).number_format = '#,##0.00'
            ws.cell(row=row, column=6, value=float(outstanding)).number_format = '#,##0.00'
            ws.cell(row=row, column=7, value=rate).number_format = '0.00'
            
            row += 1
        
        for col in range(1, 8):
            ws.column_dimensions[get_column_letter(col)].width = 18
    
    @staticmethod
    def _create_payment_methods_sheet(ws, payment_methods):
        """Create payment methods sheet"""
        ws['A1'] = 'Payment Method Breakdown'
        ws['A1'].font = Font(size=14, bold=True)
        ws.merge_cells('A1:D1')
        
        headers = ['Payment Method', 'Transaction Count', 'Total Amount (KES)', 'Average (KES)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        
        row = 4
        for method_name, transaction_count, total_amount in payment_methods:
            avg = float(total_amount) / transaction_count if transaction_count > 0 else 0
            
            ws.cell(row=row, column=1, value=method_name)
            ws.cell(row=row, column=2, value=transaction_count)
            ws.cell(row=row, column=3, value=float(total_amount)).number_format = '#,##0.00'
            ws.cell(row=row, column=4, value=avg).number_format = '#,##0.00'
            
            row += 1
        
        for col in range(1, 5):
            ws.column_dimensions[get_column_letter(col)].width = 20
    
    @staticmethod
    def _create_defaulters_sheet(ws, defaulters):
        """Create defaulters sheet"""
        ws['A1'] = 'Top Defaulters List'
        ws['A1'].font = Font(size=14, bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        ws.merge_cells('A1:G1')
        
        headers = ['Student Name', 'Admission No', 'Grade', 'Stream', 'Total Fees (KES)', 'Paid (KES)', 'Outstanding (KES)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        
        row = 4
        for student, total_fees, total_paid, balance in defaulters:
            ws.cell(row=row, column=1, value=student.name)
            ws.cell(row=row, column=2, value=student.admission_number)
            ws.cell(row=row, column=3, value=student.grade.name if student.grade else 'N/A')
            ws.cell(row=row, column=4, value=student.stream.name if student.stream else 'N/A')
            ws.cell(row=row, column=5, value=float(total_fees)).number_format = '#,##0.00'
            ws.cell(row=row, column=6, value=float(total_paid)).number_format = '#,##0.00'
            
            balance_cell = ws.cell(row=row, column=7, value=float(balance))
            balance_cell.number_format = '#,##0.00'
            balance_cell.font = Font(bold=True, color='C00000')
            
            row += 1
        
        for col in range(1, 8):
            ws.column_dimensions[get_column_letter(col)].width = 18
    
    @staticmethod
    def export_defaulters_to_pdf(defaulters_data, school_name="Hillview School"):
        """
        Export defaulters list to PDF
        
        Args:
            defaulters_data: List of defaulter records
            school_name: Name of the school
            
        Returns:
            BytesIO: PDF file as bytes
        """
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#C00000'),
            spaceAfter=30,
            alignment=1  # Center
        )
        
        # Title
        elements.append(Paragraph(f"{school_name}", styles['Heading1']))
        elements.append(Paragraph("DEFAULTERS LIST", title_style))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Summary
        total_outstanding = sum(float(balance) for _, _, _, balance in defaulters_data)
        summary_text = f"<b>Total Defaulters:</b> {len(defaulters_data)} | <b>Total Outstanding:</b> KES {total_outstanding:,.2f}"
        elements.append(Paragraph(summary_text, styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Table data
        table_data = [
            ['#', 'Student Name', 'Adm No', 'Grade', 'Total Fees', 'Paid', 'Outstanding']
        ]
        
        for idx, (student, total_fees, total_paid, balance) in enumerate(defaulters_data, 1):
            table_data.append([
                str(idx),
                student.name[:25],  # Truncate long names
                student.admission_number,
                student.grade.name if student.grade else 'N/A',
                f"{float(total_fees):,.0f}",
                f"{float(total_paid):,.0f}",
                f"{float(balance):,.0f}"
            ])
        
        # Create table
        table = Table(table_data, colWidths=[0.5*inch, 2*inch, 1*inch, 1*inch, 1*inch, 1*inch, 1.2*inch])
        
        # Table style
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5B9BD5')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            # Highlight outstanding column
            ('TEXTCOLOR', (6, 1), (6, -1), colors.HexColor('#C00000')),
            ('FONTNAME', (6, 1), (6, -1), 'Helvetica-Bold'),
        ]))
        
        elements.append(table)
        
        # Footer
        elements.append(Spacer(1, 30))
        footer_text = "This is a computer-generated document. No signature required."
        elements.append(Paragraph(footer_text, styles['Italic']))
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        
        return buffer
    
    @staticmethod
    def export_balance_report_to_excel(students_with_balances):
        """Export balance report to Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Balance Report"
        
        # Title
        ws['A1'] = 'Student Fee Balance Report'
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:H1')
        
        ws['A2'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}'
        ws.merge_cells('A2:H2')
        
        # Headers
        headers = ['Student Name', 'Admission No', 'Grade', 'Stream', 'Total Fees', 'Paid', 'Balance', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
        
        # Data
        row = 5
        for student_data in students_with_balances:
            ws.cell(row=row, column=1, value=student_data['name'])
            ws.cell(row=row, column=2, value=student_data['admission_number'])
            ws.cell(row=row, column=3, value=student_data['grade'])
            ws.cell(row=row, column=4, value=student_data['stream'])
            ws.cell(row=row, column=5, value=float(student_data['total_fees'])).number_format = '#,##0.00'
            ws.cell(row=row, column=6, value=float(student_data['total_paid'])).number_format = '#,##0.00'
            ws.cell(row=row, column=7, value=float(student_data['balance'])).number_format = '#,##0.00'
            
            status = student_data['status']
            status_cell = ws.cell(row=row, column=8, value=status)
            if status == 'Paid':
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif status == 'Partial':
                status_cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            else:
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            
            row += 1
        
        # Adjust widths
        for col in range(1, 9):
            ws.column_dimensions[get_column_letter(col)].width = 18
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
